"""Mediciones para pedir a proveedores (carpintería y hierro) a partir del modelo v8. Genera pedido_proveedores.xlsx con columnas de precio para rellenar."""
import math, collections
import casa_v8 as c
from perfiles import SECC
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ---------------- carpintería ----------------
ven, pue = [], []
for pl, ori, pos, t, ini, w, tipo, sw, tag, sill, hh in c.OPEN:
    if tipo == "L" or (t < 0.05 and tag != "PT"): continue
    ubic = {"PB": "Planta baja", "PA": "Planta alta"}[pl]
    if tag == "CR": ven.append(("Cristalera del salón (norte)", ubic, w, hh, 0.0, 1, "Corredera elevable PVC + fijo, vidrio 6/16/44.2 bajo emisivo"))
    elif tipo == "V" and t >= 0.2: ven.append(("Ventana exterior", ubic, w, hh, sill, 1, "PVC oscilobatiente 2 hojas, 4/16/4 bajo emisivo, persiana + contraventana"))
    elif tipo == "V": ven.append(("Cristal fijo interior (despacho-salón)", ubic, w, hh, sill, 1, "Fijo con perfil PVC, vidrio laminado 33.1"))
    elif tag == "PT": pue.append(("Corredera a la terraza", ubic, w, hh, 1, "PVC corredera 2 hojas, vidrio 4/16/4 bajo emisivo"))
    elif tag.startswith("P0"): pue.append(("Puerta exterior" + (" de entrada (2 hojas)" if w > 1.2 else " de la cocina al patio"), ubic, w, hh, 1, "PVC o hierro con vidrio, cerradura 3 puntos"))
    else: pue.append(("Puerta interior" + (" corredera empotrada" if tag == "P3" and w == 0.80 and sw == 0 else ""), ubic, w, hh, 1, "Hoja lisa lacada, marco de madera"))
ven.append(("Cierre acristalado de la sala familiar hacia la terraza", "Planta alta", round((c.XI1 - 2.95) - 1.42, 2), round(c.H_PA - c.T_LOSA, 2), 0.0, 1, "Fijo de suelo a techo en PVC, vidrio 4/16/4 bajo emisivo"))
# ---------------- hierro ----------------
corte = collections.OrderedDict()
def anota(cat, sec, L, cant=1):
    corte.setdefault((cat, sec, round(L, 2)), 0); corte[(cat, sec, round(L, 2))] += cant
for cat, lst in c.STEEL.items():
    for g in lst:
        if cat == "pilares": anota("Pilares", g[2], c.zt2 - c.hb)
        else:
            x0, y0, x1, y1, sec = g; L = math.hypot(x1 - x0, y1 - y0)
            anota({"vigas_forjado": "Vigas de forjado PB/PA", "vigas_cubierta": "Vigas de cubierta", "dintel": "Dintel de la cristalera del salón", "dinteles": "Dinteles de huecos"}[cat], sec, L)
rows_h = []
for (cat, sec, L), n in corte.items():
    kgm = SECC[sec][7]; rows_h.append((cat, sec, L, n, kgm, round(L * n * kgm, 1)))
extras = [("Placas base y cabeza de pilares + anclajes (12 uds, estimado)", "", "", 12, "", 12 * 14.0), ("Escalera de acero: 2 zancas UPN + 17 peldaños + descansillo (estimado)", "", "", 1, "", 380.0),
          ("Barandillas: lado abierto de la escalera, suelo junto al hueco y pasamanos (≈ 10 m, estimado)", "", "", 10, "", 10 * 14.0), (f"Soportes de {len(c.PANELES)} paneles solares, acero galvanizado (≈ 5 kg/panel, estimado)", "", "", len(c.PANELES), "", len(c.PANELES) * 5.0),
          ("Tornillería, cartelas y piezas de unión (≈ 3 %)", "", "", 1, "", round(0.03 * sum(r[5] for r in rows_h), 1))]
kg_tot = sum(r[5] for r in rows_h) + sum(e[5] for e in extras)

wb = Workbook(); azul = PatternFill("solid", fgColor="DDEBF7"); amar = PatternFill("solid", fgColor="FFF2CC"); neg = Font(bold=True)
def cab(ws, fila, textos):
    for i, t in enumerate(textos, 1):
        cl = ws.cell(row=fila, column=i, value=t); cl.font = neg; cl.fill = azul; cl.alignment = Alignment(wrap_text=True, vertical="center")
ws = wb.active; ws.title = "Ventanas"
ws["A1"] = "Ventanas, cristaleras y fijos (medidas del modelo v8, huecos de obra: ancho × alto en m)"; ws["A1"].font = Font(bold=True, size=12)
cab(ws, 3, ["Elemento", "Planta", "Ancho (m)", "Alto (m)", "Antepecho (m)", "Uds", "Superficie total (m²)", "Descripción", "Precio unidad (€) – RELLENAR", "Total (€)"])
for i, (n, p, w, h, s, u, d) in enumerate(ven, 4):
    ws.append([n, p, w, h, s, u, f"=C{i}*D{i}*F{i}", d, None, f"=I{i}*F{i}"]); ws.cell(row=i, column=9).fill = amar
f = 4 + len(ven)
ws.cell(row=f, column=1, value="TOTAL ventanas").font = neg; ws.cell(row=f, column=7, value=f"=SUM(G4:G{f-1})"); ws.cell(row=f, column=10, value=f"=SUM(J4:J{f-1})").font = neg
ws.cell(row=f + 2, column=1, value="Contraventanas de lamas: 2 hojas por cada ventana exterior (fachada sur y norte). Mosquiteras y persianas: según tarifa.")
ws2 = wb.create_sheet("Puertas"); ws2["A1"] = "Puertas (huecos de obra)"; ws2["A1"].font = Font(bold=True, size=12)
cab(ws2, 3, ["Elemento", "Planta", "Ancho (m)", "Alto (m)", "Uds", "Descripción", "Precio unidad (€) – RELLENAR", "Total (€)"])
for i, (n, p, w, h, u, d) in enumerate(pue, 4):
    ws2.append([n, p, w, h, u, d, None, f"=G{i}*E{i}"]); ws2.cell(row=i, column=7).fill = amar
f2 = 4 + len(pue); ws2.cell(row=f2, column=1, value="TOTAL puertas").font = neg; ws2.cell(row=f2, column=8, value=f"=SUM(H4:H{f2-1})").font = neg
ws3 = wb.create_sheet("Hierro"); ws3["A1"] = "Lista de corte y pedido de hierro (S275JR). Longitudes entre ejes de apoyo; añadir 1–2 % de merma de corte"; ws3["A1"].font = Font(bold=True, size=12)
cab(ws3, 3, ["Pieza", "Perfil", "Longitud (m)", "Uds", "kg/m", "Peso (kg)", "Precio €/kg – RELLENAR", "Total (€)"])
r0 = 4
for i, r in enumerate(rows_h + [e for e in extras], r0):
    cat, sec, L, n, kgm, kg = r
    ws3.append([cat, sec, L, n, kgm, kg, None, f"=F{i}*G{i}"]); ws3.cell(row=i, column=7).fill = amar
f3 = r0 + len(rows_h) + len(extras)
ws3.cell(row=f3, column=1, value="TOTAL hierro").font = neg; ws3.cell(row=f3, column=6, value=f"=SUM(F{r0}:F{f3-1})").font = neg; ws3.cell(row=f3, column=8, value=f"=SUM(H{r0}:H{f3-1})").font = neg
ws3.cell(row=f3 + 2, column=1, value="Las filas sin perfil son estimaciones mías (escalera, barandillas, soportes solares, placas): confirmar con el taller. Las secciones de las vigas son una propuesta que debe calcular un ingeniero.")
ws4 = wb.create_sheet("Resumen y ahorro"); ws4["A1"] = "Resumen y ahorro frente a precio de mercado"; ws4["A1"].font = Font(bold=True, size=12)
cab(ws4, 3, ["Concepto", "Tu coste (€)", "Referencia mercado (€) – RELLENAR", "Ahorro (€)", "Ahorro (%)"])
ws4.append(["Ventanas y cristaleras", f"=Ventanas!J{f}", None, "=C4-B4", "=IF(C4=0,\"\",D4/C4)"])
ws4.append(["Puertas", f"=Puertas!H{f2}", None, "=C5-B5", "=IF(C5=0,\"\",D5/C5)"])
ws4.append(["Hierro (estructura, escalera, barandillas)", f"=Hierro!H{f3}", None, "=C6-B6", "=IF(C6=0,\"\",D6/C6)"])
ws4.append(["TOTAL", "=SUM(B4:B6)", "=SUM(C4:C6)", "=C7-B7", "=IF(C7=0,\"\",D7/C7)"])
for r in (4, 5, 6): ws4.cell(row=r, column=3).fill = amar
for r in range(4, 8): ws4.cell(row=r, column=5).number_format = "0%"
for w_ in (ws, ws2, ws3, ws4):
    for col in range(1, 11): w_.column_dimensions[get_column_letter(col)].width = 30 if col in (1, 8) else 16
wb.save("pedido_proveedores.xlsx")
print("ventanas:", len(ven), "puertas:", len(pue), "m² ventanas:", round(sum(w * h * u for _, _, w, h, _, u, _ in ven), 1), "hierro kg:", round(kg_tot))
for r in rows_h: print(r)
